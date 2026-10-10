"""Train clinical models on train and compare them on combined validation data."""

import argparse
import csv
import json
import pickle
import time
from pathlib import Path

import numpy as np
from lightgbm import LGBMClassifier
from openpyxl import load_workbook
from sklearn.calibration import calibration_curve
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import ExtraTreesClassifier, RandomForestClassifier
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    average_precision_score,
    brier_score_loss,
    confusion_matrix,
    roc_auc_score,
    roc_curve,
)
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.svm import SVC
from xgboost import XGBClassifier

from Clinical_data_screening import (
    CATEGORICAL_VARIABLES,
    NUMERIC_VARIABLES,
    parse_category,
    parse_number,
)


def read_rows(workbook_path, preprocess_root):
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    index = {name: i + 1 for i, name in enumerate(headers)}
    split_map = {}
    for split in ("train", "validation"):
        split_dir = preprocess_root / split
        if not split_dir.exists():
            raise FileNotFoundError(f"缺少数据集文件夹: {split_dir}")
        for patient_dir in split_dir.iterdir():
            if patient_dir.is_dir() and patient_dir.name.isdigit():
                if patient_dir.name in split_map:
                    raise ValueError(f"患者重复出现在多个数据集: {patient_dir.name}")
                split_map[patient_dir.name] = split
    rows = []
    for row_number in range(2, sheet.max_row + 1):
        patient_id = sheet.cell(row_number, index["编号"]).value
        patient_id = (
            str(int(patient_id))
            if isinstance(patient_id, float) and patient_id.is_integer()
            else str(patient_id).strip()
        )
        evaluation = str(sheet.cell(row_number, index["Evaluation"]).value).strip()
        if evaluation not in {"1", "2"} or patient_id not in split_map:
            continue
        row = {"编号": patient_id, "y": 1 if evaluation == "1" else 0, "split": split_map[patient_id]}
        for name in NUMERIC_VARIABLES:
            row[name] = parse_number(sheet.cell(row_number, index[name]).value)
        for name in CATEGORICAL_VARIABLES:
            row[name] = parse_category(sheet.cell(row_number, index[name]).value)
        rows.append(row)
    return rows


def build_matrix(rows, variables):
    return np.array([[row[variable] for variable in variables] for row in rows], dtype=object)


def make_preprocessor(variables, scale_numeric):
    numeric_indices = [i for i, name in enumerate(variables) if name in NUMERIC_VARIABLES]
    categorical_indices = [i for i, name in enumerate(variables) if name in CATEGORICAL_VARIABLES]
    numeric_steps = [("imputer", SimpleImputer(strategy="median"))]
    if scale_numeric:
        numeric_steps.append(("scaler", StandardScaler()))
    return ColumnTransformer([
        ("numeric", Pipeline(numeric_steps), numeric_indices),
        ("categorical", Pipeline([
            ("imputer", SimpleImputer(strategy="most_frequent")),
            ("onehot", OneHotEncoder(handle_unknown="ignore")),
        ]), categorical_indices),
    ], remainder="drop")


def positive_probability(model, x):
    if hasattr(model, "predict_proba"):
        return model.predict_proba(x)[:, 1]
    return model.decision_function(x)


def metrics(y_true, probability, threshold=0.5):
    prediction = (np.asarray(probability) >= threshold).astype(int)
    tn, fp, fn, tp = confusion_matrix(y_true, prediction, labels=[0, 1]).ravel()
    return {
        "roc_auc": float(roc_auc_score(y_true, probability)),
        "pr_auc": float(average_precision_score(y_true, probability)),
        "brier_score": float(brier_score_loss(y_true, probability)),
        "sensitivity": float(tp / (tp + fn)) if tp + fn else float("nan"),
        "specificity": float(tn / (tn + fp)) if tn + fp else float("nan"),
        "accuracy": float(accuracy_score(y_true, prediction)),
        "threshold": float(threshold),
        "tn": int(tn), "fp": int(fp), "fn": int(fn), "tp": int(tp),
    }


def select_threshold(y_true, probability):
    fpr, tpr, thresholds = roc_curve(y_true, probability)
    valid = np.isfinite(thresholds)
    youden = tpr[valid] - fpr[valid]
    valid_thresholds = thresholds[valid]
    best = np.flatnonzero(youden == np.max(youden))
    return float(valid_thresholds[best[np.argmin(np.abs(valid_thresholds[best] - 0.5))]])


def bootstrap_ci(y_true, probability, threshold, repeats, rng):
    names = ("roc_auc", "pr_auc", "brier_score", "sensitivity", "specificity", "accuracy")
    values = {name: [] for name in names}
    y_true = np.asarray(y_true)
    probability = np.asarray(probability)
    for _ in range(repeats):
        indices = rng.integers(0, len(y_true), size=len(y_true))
        if len(np.unique(y_true[indices])) < 2:
            continue
        result = metrics(y_true[indices], probability[indices], threshold)
        for name in names:
            if np.isfinite(result[name]):
                values[name].append(result[name])
    output = {}
    for name, samples in values.items():
        output[f"{name}_ci_lower"] = float(np.percentile(samples, 2.5)) if samples else float("nan")
        output[f"{name}_ci_upper"] = float(np.percentile(samples, 97.5)) if samples else float("nan")
    return output


def calibration_metrics(y_true, probability):
    clipped = np.clip(probability, 1e-6, 1 - 1e-6)
    logit = np.log(clipped / (1 - clipped))
    calibration_model = LogisticRegression(penalty=None, solver="lbfgs", max_iter=2000)
    calibration_model.fit(logit.reshape(-1, 1), y_true)
    intercept = 0.0
    for _ in range(100):
        fitted = 1.0 / (1.0 + np.exp(-np.clip(logit + intercept, -30, 30)))
        step = np.sum(y_true - fitted) / max(np.sum(fitted * (1 - fitted)), 1e-12)
        intercept += step
        if abs(step) < 1e-10:
            break
    return {
        "calibration_intercept": float(intercept),
        "calibration_slope": float(calibration_model.coef_[0, 0]),
    }


def calibration_curve_rows(y_true, probability, n_bins=10):
    fraction, mean_predicted = calibration_curve(
        y_true, probability, n_bins=n_bins, strategy="quantile"
    )
    edges = np.unique(np.quantile(probability, np.linspace(0, 1, n_bins + 1)))
    bin_ids = np.digitize(probability, edges[1:-1], right=True)
    return [
        {
            "split": "validation_combined",
            "bin": index + 1,
            "mean_predicted_probability": float(predicted),
            "observed_CR_fraction": float(observed),
            "n": int(np.sum(bin_ids == index)),
        }
        for index, (observed, predicted) in enumerate(zip(fraction, mean_predicted))
    ]


def plot_calibration_curve(curve_rows, output_path):
    try:
        import matplotlib.pyplot as plt
    except ImportError:
        return False
    figure, axis = plt.subplots(figsize=(5, 5))
    axis.plot([0, 1], [0, 1], "--", color="gray", label="Perfect calibration")
    axis.plot(
        [row["mean_predicted_probability"] for row in curve_rows],
        [row["observed_CR_fraction"] for row in curve_rows],
        marker="o", label="validation_combined",
    )
    axis.set_xlabel("Predicted probability of CR")
    axis.set_ylabel("Observed CR fraction")
    axis.set_title("Clinical model calibration")
    axis.legend()
    figure.tight_layout()
    figure.savefig(output_path, dpi=300)
    plt.close(figure)
    return True


def model_specs(random_state):
    return {
        "LogisticRegression": (LogisticRegression(
            max_iter=3000, class_weight="balanced", solver="liblinear", random_state=random_state
        ), True),
        "SVM": (SVC(probability=True, class_weight="balanced", random_state=random_state), True),
        "RandomForest": (RandomForestClassifier(
            n_estimators=500, class_weight="balanced", n_jobs=1, random_state=random_state
        ), False),
        "ExtraTrees": (ExtraTreesClassifier(
            n_estimators=500, class_weight="balanced", n_jobs=1, random_state=random_state
        ), False),
        "LightGBM": (LGBMClassifier(
            n_estimators=300, learning_rate=0.03, verbosity=-1, random_state=random_state, n_jobs=1
        ), False),
        "XGBoost": (XGBClassifier(
            n_estimators=300, max_depth=3, learning_rate=0.03,
            subsample=0.8, colsample_bytree=0.8, eval_metric="logloss",
            tree_method="hist", random_state=random_state, n_jobs=1,
        ), False),
    }


def make_pipeline(name, variables, random_state):
    estimator, scale_numeric = model_specs(random_state)[name]
    return Pipeline([
        ("preprocessor", make_preprocessor(variables, scale_numeric)),
        ("model", estimator),
    ])


def write_csv(path, rows, fieldnames=None):
    if fieldnames is None:
        fieldnames = sorted({key for row in rows for key in row}) if rows else []
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def leave_one_out_training_influence(train, variables, best_name, random_state):
    start_time = time.perf_counter()
    x_train = build_matrix(train, variables)
    y_train = np.array([row["y"] for row in train])
    full_pipeline = make_pipeline(best_name, variables, random_state)
    full_pipeline.fit(x_train, y_train)
    full_probability = positive_probability(full_pipeline, x_train)
    full_auc = roc_auc_score(y_train, full_probability)
    results = []
    for index, row in enumerate(train, start=1):
        row_index = index - 1
        keep = np.ones(len(train), dtype=bool)
        keep[row_index] = False
        loo_pipeline = make_pipeline(best_name, variables, random_state)
        loo_pipeline.fit(x_train[keep], y_train[keep])
        reduced_probability = positive_probability(loo_pipeline, x_train[keep])
        auc_without = roc_auc_score(y_train[keep], reduced_probability)
        results.append({
            "编号": row["编号"],
            "y_true": row["y"],
            "full_model_CR_probability": float(full_probability[row_index]),
            "full_training_auc": float(full_auc),
            "auc_without_patient": float(auc_without),
            "delta_auc_after_removal": float(auc_without - full_auc),
        })
        if index == 1 or index == len(train) or index % 25 == 0:
            elapsed = time.perf_counter() - start_time
            eta = elapsed / index * (len(train) - index)
            print(f"[留一法 {index}/{len(train)}] 已用 {elapsed:.1f} 秒，预计剩余 {eta:.1f} 秒")
    return results


def main():
    start_time = time.perf_counter()
    print("[1/4] 开始读取数据，预计耗时：数分钟")
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=Path("Clinical_data_cleaning.xlsx"))
    parser.add_argument("--preprocess-root", type=Path, default=Path("Preprocess"))
    parser.add_argument("--screening-dir", type=Path, default=Path("outputs/final/01_临床筛选"))
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/final/02_模型评估"))
    parser.add_argument("--patient-output-dir", type=Path, default=Path("outputs/final/03_患者级结果"))
    parser.add_argument("--random-state", type=int, default=None)
    parser.add_argument("--bootstrap-repeats", type=int, default=2000)
    args = parser.parse_args()

    rows = read_rows(args.workbook, args.preprocess_root)
    train = [row for row in rows if row["split"] == "train"]
    validation = [row for row in rows if row["split"] == "validation"]
    selected_path = args.screening_dir / "selected_predictors.json"
    variables = []
    if selected_path.exists():
        selected = json.loads(selected_path.read_text(encoding="utf-8"))
        variables = selected.get("independent_predictors") or selected.get("univariable_p_lt_0_05", [])
    if not variables:
        raise ValueError("没有可用于机器学习的临床变量")

    x_train = build_matrix(train, variables)
    y_train = np.array([row["y"] for row in train])
    x_validation = build_matrix(validation, variables)
    y_validation = np.array([row["y"] for row in validation])
    print(f"[2/4] 已读取训练集 {len(train)} 例、合并验证集 {len(validation)} 例；开始拟合6个模型")
    comparison = []
    fitted_models = {}
    validation_probabilities = {}
    model_names = list(model_specs(args.random_state))
    for index, name in enumerate(model_names, start=1):
        model_start = time.perf_counter()
        print(f"[模型 {index}/{len(model_names)}] 开始拟合 {name}")
        pipeline = make_pipeline(name, variables, args.random_state)
        pipeline.fit(x_train, y_train)
        probability = positive_probability(pipeline, x_validation)
        fitted_models[name] = pipeline
        validation_probabilities[name] = probability
        comparison.append({"model": name, **metrics(y_validation, probability, 0.5)})
        elapsed = time.perf_counter() - start_time
        eta = elapsed / index * (len(model_names) - index)
        print(f"[模型 {index}/{len(model_names)}] {name} 完成，本模型耗时 {time.perf_counter() - model_start:.1f} 秒；预计剩余 {eta:.1f} 秒")

    comparison.sort(key=lambda row: (row["roc_auc"], row["pr_auc"]), reverse=True)
    best_name = comparison[0]["model"]
    best_pipeline = fitted_models[best_name]
    validation_probability = validation_probabilities[best_name]
    threshold = select_threshold(y_validation, validation_probability)
    rng = np.random.default_rng(args.random_state)
    holdout = {
        "split": "validation_combined",
        "model": best_name,
        **metrics(y_validation, validation_probability, threshold),
        **bootstrap_ci(y_validation, validation_probability, threshold, args.bootstrap_repeats, rng),
        **calibration_metrics(y_validation, validation_probability),
    }
    print(f"[3/4] 最佳模型为 {best_name}，开始训练集留一法影响分析，预计耗时：数分钟")
    loo_influence = leave_one_out_training_influence(
        train, variables, best_name, args.random_state
    )

    patient_predictions = []
    for split_name, split_rows, probability in (
        ("train", train, positive_probability(best_pipeline, x_train)),
        ("validation_combined", validation, validation_probability),
    ):
        clipped = np.clip(probability, 1e-6, 1 - 1e-6)
        for row, probability_value in zip(split_rows, clipped):
            patient_predictions.append({
                "编号": row["编号"],
                "split": split_name,
                "y_true": row["y"],
                "CR_probability": float(probability_value),
                "clinical_score_logit": float(np.log(probability_value / (1 - probability_value))),
                "decision_threshold": float(threshold),
                "predicted_y_at_validation_threshold": int(probability_value >= threshold),
            })

    curve_rows = calibration_curve_rows(y_validation, validation_probability)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.patient_output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "validation_model_comparison.csv", comparison)
    write_csv(args.output_dir / "best_model_holdout_metrics.csv", [holdout])
    write_csv(args.output_dir / "calibration_curve.csv", curve_rows)
    write_csv(args.patient_output_dir / "training_leave_one_out_influence.csv", loo_influence)
    write_csv(args.patient_output_dir / "clinical_patient_predictions.csv", patient_predictions)
    plot_created = plot_calibration_curve(curve_rows, args.output_dir / "calibration_curve.png")
    (args.output_dir / "clinical_best_model.json").write_text(json.dumps({
        "model": best_name,
        "variables": variables,
        "random_state": args.random_state,
        "validation_strategy": "fit_all_models_on_train_compare_on_combined_validation",
        "bootstrap_repeats": args.bootstrap_repeats,
        "validation_selected_threshold": threshold,
        "calibration_plot_created": plot_created,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    with (args.output_dir / "best_clinical_model.pkl").open("wb") as model_file:
        pickle.dump(best_pipeline, model_file)
    print(f"[4/4] 建模完成，结果已写入：{args.output_dir}；患者结果：{args.patient_output_dir}")
    print(f"总耗时: {time.perf_counter() - start_time:.1f} 秒")
    print(json.dumps({
        "rows": len(rows),
        "train_rows": len(train),
        "validation_rows": len(validation),
        "best_model": best_name,
        "variables": variables,
        "validation_model_comparison": comparison,
        "validation_threshold": threshold,
        "holdout_metrics": [holdout],
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
