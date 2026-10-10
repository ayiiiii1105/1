"""Clinical-variable screening for CR vs non-CR.

CR is encoded as y=1 and non-CR as y=0. Screening is performed on the
folder-defined training cohort only.
"""

import argparse
import csv
import json
import math
import time
from pathlib import Path

import numpy as np
from scipy.stats import chi2, norm

from openpyxl import load_workbook


NUMERIC_VARIABLES = [
    "Age(Years)", "Height(cm)", "Weight(kg)", "BMI",
    "WBC(10^9/L)", "Lymphocyte(10^9/L)", "Monocyte(10^9/L)",
    "Neutrophils(10^9/L)", "PLT(10^9/L)", "NLR", "MLR", "PLR",
    "ALP(U/L)", "LDH(U/L)", "Alb(g/L)",
    "Thevolumeofnasopharyngealtumor（mm3）",
]
CATEGORICAL_VARIABLES = [
    "Gender", "Smoking", "Drinking", "FamilyhistoryofNPC",
    "Tumorinvasion(T-stage)", "Lymphnodemetastasis(N-stage)",
    "Clinicalstage",
]
GTV_VOLUME_VARIABLE = "Thevolumeofnasopharyngealtumor（mm3）"


def parse_number(value):
    if value is None or str(value).strip() == "":
        return np.nan
    try:
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def parse_category(value):
    if value is None or str(value).strip() == "":
        return None
    return str(value).strip()


def read_data(workbook_path, preprocess_root):
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    index = {name: i + 1 for i, name in enumerate(headers)}
    required = ["编号", "Evaluation", *NUMERIC_VARIABLES, *CATEGORICAL_VARIABLES]
    missing = [name for name in required if name not in index]
    if missing:
        raise ValueError(f"缺少字段: {missing}")

    split_map = {}
    for split in ("train", "validation"):
        split_dir = preprocess_root / split
        if not split_dir.exists():
            raise FileNotFoundError(f"缺少数据集文件夹: {split_dir}")
        for patient_dir in split_dir.iterdir():
            if patient_dir.is_dir() and patient_dir.name.isdigit():
                if patient_dir.name in split_map:
                    raise ValueError(f"患者重复出现在多个数据集: {patient_dir.name}")
                split_map[patient_dir.name] = split

    rows = []
    for row_number in range(2, sheet.max_row + 1):
        patient_id = sheet.cell(row_number, index["编号"]).value
        patient_id = str(int(patient_id)) if isinstance(patient_id, float) and patient_id.is_integer() else str(patient_id).strip()
        evaluation = str(sheet.cell(row_number, index["Evaluation"]).value).strip()
        if evaluation not in {"1", "2"} or patient_id not in split_map:
            continue
        row = {"编号": patient_id, "Evaluation": evaluation, "split": split_map[patient_id]}
        for name in NUMERIC_VARIABLES:
            row[name] = parse_number(sheet.cell(row_number, index[name]).value)
        for name in CATEGORICAL_VARIABLES:
            row[name] = parse_category(sheet.cell(row_number, index[name]).value)
        row["y"] = 1 if evaluation == "1" else 0
        rows.append(row)
    return rows


def matrix_for(rows, variables):
    complete = []
    for row in rows:
        if all((not np.isnan(row[name])) if name in NUMERIC_VARIABLES else row[name] is not None for name in variables):
            complete.append(row)
    if not complete:
        raise ValueError(f"变量没有完整记录: {variables}")
    columns = ["const"]
    matrix = [np.ones(len(complete))]
    for name in variables:
        if name in NUMERIC_VARIABLES:
            columns.append(name)
            matrix.append(np.array([row[name] for row in complete], dtype=float))
        else:
            levels = sorted({row[name] for row in complete})
            for level in levels[1:]:
                columns.append(f"{name}={level}")
                matrix.append(np.array([1.0 if row[name] == level else 0.0 for row in complete]))
    return complete, np.column_stack(matrix), columns


def impute_rows(rows, variables):
    imputation = {}
    imputed_rows = [dict(row) for row in rows]
    for variable in variables:
        if variable in NUMERIC_VARIABLES:
            observed = [row[variable] for row in rows if not np.isnan(row[variable])]
            if not observed:
                raise ValueError(f"变量全部缺失，无法插补: {variable}")
            if variable == GTV_VOLUME_VARIABLE and len(observed) != len(rows):
                raise ValueError("GTV 体积存在缺失；请先从 structures.nii.gz 计算 GTV 体积，禁止中位数填补")
            value = float(np.median(observed))
            imputation[variable] = {"method": "median", "value": value, "missing_n": len(rows) - len(observed)}
        else:
            observed = [row[variable] for row in rows if row[variable] is not None]
            if not observed:
                raise ValueError(f"变量全部缺失，无法插补: {variable}")
            values, counts = np.unique(observed, return_counts=True)
            value = str(values[int(np.argmax(counts))])
            imputation[variable] = {"method": "mode", "value": value, "missing_n": len(rows) - len(observed)}
        for row in imputed_rows:
            if variable in NUMERIC_VARIABLES:
                if np.isnan(row[variable]):
                    row[variable] = value
            elif row[variable] is None:
                row[variable] = value
    return imputed_rows, imputation


def firth_logit(x, y, max_iter=200, tolerance=1e-8):
    beta = np.zeros(x.shape[1], dtype=float)
    for _ in range(max_iter):
        eta = np.clip(x @ beta, -30, 30)
        p = 1.0 / (1.0 + np.exp(-eta))
        w = np.maximum(p * (1 - p), 1e-8)
        xtwx = x.T @ (w[:, None] * x)
        cov = np.linalg.pinv(xtwx)
        hat = np.sum((x @ cov) * x, axis=1) * w
        adjusted_y = y - p + hat * (0.5 - p)
        step = cov @ (x.T @ adjusted_y)
        new_beta = beta + step
        if np.max(np.abs(step)) < tolerance:
            beta = new_beta
            break
        beta = new_beta
    eta = np.clip(x @ beta, -30, 30)
    p = 1.0 / (1.0 + np.exp(-eta))
    w = np.maximum(p * (1 - p), 1e-8)
    cov = np.linalg.pinv(x.T @ (w[:, None] * x))
    se = np.sqrt(np.maximum(np.diag(cov), 1e-12))
    z = beta / se
    p_values = 2 * norm.sf(np.abs(z))
    return beta, se, p_values


def standard_logit(x, y, max_iter=200, tolerance=1e-8):
    beta = np.zeros(x.shape[1], dtype=float)
    converged = False
    for _ in range(max_iter):
        eta = np.clip(x @ beta, -30, 30)
        p = 1.0 / (1.0 + np.exp(-eta))
        w = np.maximum(p * (1 - p), 1e-8)
        information = x.T @ (w[:, None] * x)
        step = np.linalg.pinv(information) @ (x.T @ (y - p))
        beta += step
        if np.max(np.abs(step)) < tolerance:
            converged = True
            break
    eta = np.clip(x @ beta, -30, 30)
    p = 1.0 / (1.0 + np.exp(-eta))
    w = np.maximum(p * (1 - p), 1e-8)
    cov = np.linalg.pinv(x.T @ (w[:, None] * x))
    se = np.sqrt(np.maximum(np.diag(cov), 1e-12))
    z = beta / se
    p_values = 2 * norm.sf(np.abs(z))
    if not converged or not np.all(np.isfinite(beta)) or np.max(np.abs(beta)) > 25:
        raise ArithmeticError("complete or quasi-complete separation")
    return beta, se, p_values


def fit_logistic(rows, variables):
    complete, x, columns = matrix_for(rows, variables)
    y = np.array([row["y"] for row in complete], dtype=float)
    method = "standard_logistic"
    try:
        beta, se, p_values = standard_logit(x, y)
    except (ArithmeticError, np.linalg.LinAlgError, ValueError):
        method = "firth_logistic"
        beta, se, p_values = firth_logit(x, y)
    wald = float(np.sum((beta[1:] / se[1:]) ** 2)) if len(beta) > 1 else 0.0
    overall_p = float(chi2.sf(wald, max(len(beta) - 1, 1)))
    records = []
    for column, coefficient, standard_error, p_value in zip(columns[1:], beta[1:], se[1:], p_values[1:]):
        predictor_name = column.split("=", 1)[0]
        records.append({
            "predictor": predictor_name,
            "term": column,
            "n": len(complete),
            "method": method,
            "coefficient": float(coefficient),
            "odds_ratio": float(math.exp(np.clip(coefficient, -50, 50))),
            "p_value": float(p_value),
            "overall_p_value": overall_p,
        })
    return records, overall_p, method, len(complete)


def calculate_vif(rows, variables):
    complete, x, columns = matrix_for(rows, variables)
    if x.shape[1] <= 2:
        return []
    values = []
    for column_index in range(1, x.shape[1]):
        target = x[:, column_index]
        others = np.delete(x, column_index, axis=1)
        coef = np.linalg.pinv(others.T @ others) @ others.T @ target
        predicted = others @ coef
        ss_res = np.sum((target - predicted) ** 2)
        ss_tot = np.sum((target - np.mean(target)) ** 2)
        r_squared = 0.0 if ss_tot <= 1e-12 else max(0.0, min(1.0, 1 - ss_res / ss_tot))
        values.append({"term": columns[column_index], "vif": float(1 / max(1 - r_squared, 1e-12)), "n": len(complete)})
    return values


def write_csv(path, rows, fieldnames):
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    start_time = time.perf_counter()
    print("[1/3] 开始读取临床数据，预计耗时：少于1分钟")
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=Path("Clinical_data_cleaning.xlsx"))
    parser.add_argument("--preprocess-root", type=Path, default=Path("Preprocess"))
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/final/01_临床筛选"))
    args = parser.parse_args()
    rows = read_data(args.workbook, args.preprocess_root)
    train_rows = [row for row in rows if row["split"] == "train"]
    print(f"[2/3] 已读取 {len(rows)} 例，其中训练集 {len(train_rows)} 例；开始单因素 Logistic 筛选")
    univariate = []
    selected = []
    variables_to_screen = [*NUMERIC_VARIABLES, *CATEGORICAL_VARIABLES]
    for index, variable in enumerate(variables_to_screen, start=1):
        result, overall_p, _, _ = fit_logistic(train_rows, [variable])
        univariate.extend(result)
        if overall_p < 0.05:
            selected.append(variable)
        if index == 1 or index == len(variables_to_screen) or index % 5 == 0:
            elapsed = time.perf_counter() - start_time
            eta = elapsed / index * (len(variables_to_screen) - index)
            print(f"[筛选 {index}/{len(variables_to_screen)}] 已用 {elapsed:.1f} 秒，预计剩余 {eta:.1f} 秒")
    multivariate = []
    if selected:
        imputed_train_rows, imputation = impute_rows(train_rows, selected)
        multivariate, _, method, n = fit_logistic(imputed_train_rows, selected)
        vif = calculate_vif(imputed_train_rows, selected)
    else:
        method, n, vif, imputation = "not_run", 0, [], {}
    independent = sorted({row["predictor"] for row in multivariate if row["overall_p_value"] < 0.05})
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "univariable_logistic.csv", univariate, list(univariate[0]) if univariate else ["predictor"])
    write_csv(args.output_dir / "multivariable_logistic.csv", multivariate, list(multivariate[0]) if multivariate else ["predictor"])
    write_csv(args.output_dir / "vif.csv", vif, ["term", "vif", "n"])
    with (args.output_dir / "selected_predictors.json").open("w", encoding="utf-8") as handle:
        json.dump({"univariable_p_lt_0_05": selected, "independent_predictors": independent, "multivariable_method": method, "multivariable_n": n, "imputation": imputation}, handle, ensure_ascii=False, indent=2)
    dataset_fields = ["编号", "Evaluation", "y", "split", *selected]
    write_csv(args.output_dir / "clinical_screened_dataset.csv", [{field: row.get(field) for field in dataset_fields} for row in rows], dataset_fields)
    elapsed = time.perf_counter() - start_time
    print("[3/3] 临床筛选完成，结果已写入：", args.output_dir)
    print(f"总耗时: {elapsed:.1f} 秒")
    print(json.dumps({"rows": len(rows), "train_rows": len(train_rows), "univariable_selected": selected, "independent_predictors": independent}, ensure_ascii=False))


if __name__ == "__main__":
    main()
