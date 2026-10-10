from copy import copy
from pathlib import Path
import time

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator


ROOT = Path(__file__).resolve().parent.parent
INPUT_FILE = ROOT / "Clinical_data.xlsx"
OUTPUT_FILE = ROOT / "Clinical_data_cleaning.xlsx"
SOURCE_SHEET = "Sheet1"
CLEANED_SHEET = "Cleaned_Data"


def copy_cell(source_cell, target_cell):
    if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
        target_cell.value = Translator(
            source_cell.value, origin=source_cell.coordinate
        ).translate_formula(target_cell.coordinate)
    else:
        target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format
    if source_cell.alignment:
        target_cell.alignment = copy(source_cell.alignment)
    if source_cell.protection:
        target_cell.protection = copy(source_cell.protection)


def is_valid_evaluation(value):
    return str(value).strip() in {"1", "2"}


def clean_workbook():
    start_time = time.perf_counter()
    print("[1/1] 开始清洗临床表，预计耗时：少于1分钟")
    source_workbook = load_workbook(INPUT_FILE)
    source = source_workbook[SOURCE_SHEET]
    headers = [source.cell(1, column).value for column in range(1, source.max_column + 1)]
    data_column_count = max(
        index + 1
        for index, value in enumerate(headers)
        if value is not None and str(value).strip() != ""
    )
    try:
        evaluation_column = headers.index("Evaluation") + 1
    except ValueError as exc:
        raise ValueError("未找到 Evaluation 列") from exc

    workbook = Workbook()
    cleaned = workbook.active
    cleaned.title = CLEANED_SHEET
    for column in range(1, data_column_count + 1):
        copy_cell(source.cell(1, column), cleaned.cell(1, column))

    target_row = 2
    actual_source_rows = 0
    excluded_rows = []
    for source_row in range(2, source.max_row + 1):
        row_values = [
            source.cell(source_row, column).value
            for column in range(1, data_column_count + 1)
        ]
        if not any(value is not None and str(value).strip() != "" for value in row_values):
            continue
        actual_source_rows += 1
        evaluation = source.cell(source_row, evaluation_column).value
        if not is_valid_evaluation(evaluation):
            excluded_rows.append((source_row, evaluation))
            continue
        for column in range(1, data_column_count + 1):
            copy_cell(source.cell(source_row, column), cleaned.cell(target_row, column))
        cleaned.row_dimensions[target_row].height = source.row_dimensions[source_row].height
        target_row += 1

    for column_letter, dimension in source.column_dimensions.items():
        cleaned.column_dimensions[column_letter].width = dimension.width
        cleaned.column_dimensions[column_letter].hidden = dimension.hidden
    cleaned.freeze_panes = "A2"
    cleaned.auto_filter.ref = cleaned.dimensions
    workbook.save(OUTPUT_FILE)
    elapsed = time.perf_counter() - start_time
    print(
        f"输出文件: {OUTPUT_FILE}\n"
        f"原始数据行数: {actual_source_rows}\n"
        f"清洗后数据行数: {target_row - 2}\n"
        f"排除行数: {len(excluded_rows)}\n"
        f"排除行: {excluded_rows}\n"
        f"总耗时: {elapsed:.1f} 秒"
    )


if __name__ == "__main__":
    clean_workbook()
